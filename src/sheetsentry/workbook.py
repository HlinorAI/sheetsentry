"""Internal workbook model.

openpyxl cannot hand back a formula and its computed value at the same time,
so the file is read in two passes and stitched together here. This is the
only module that knows about openpyxl; every check works against the model.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#NUM!", "#N/A"}

# Explicit "no value" markers. They are not treated as a foreign type in SD004.
BLANK_MARKERS = {"", "-", "--", "n/a", "na", "none", "null", "nil", "?", "tbd", "unknown"}

_KEY_HEADER_RE = re.compile(
    r"\b(id|code|sku|key|ean|gtin|upc|isbn|no|num|number|ref|barcode|part|serial|slug)\b",
    re.IGNORECASE,
)


@dataclass
class Cell:
    row: int
    col: int
    formula: str | None
    value: Any
    """Value cached by the last recalculation (the data_only=True pass)."""
    data_type: str
    """openpyxl type: n(number) s(string) b(bool) d(date) e(error) f(formula)."""
    number_format: str = "General"
    is_date: bool = False

    @property
    def coord(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def is_formula(self) -> bool:
        return self.formula is not None

    @property
    def is_empty(self) -> bool:
        return self.formula is None and (self.value is None or self.value == "")

    @property
    def is_numeric_constant(self) -> bool:
        return (
            not self.is_formula
            and isinstance(self.value, (int, float))
            and not isinstance(self.value, bool)
        )

    @property
    def is_text(self) -> bool:
        return not self.is_formula and isinstance(self.value, str)

    @property
    def is_blank_marker(self) -> bool:
        return self.is_text and self.value.strip().lower() in BLANK_MARKERS


@dataclass
class Sheet:
    name: str
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)
    max_row: int = 0
    max_col: int = 0
    hidden_rows: set[int] = field(default_factory=set)
    hidden_cols: set[int] = field(default_factory=set)
    merged: list[str] = field(default_factory=list)
    hidden: bool = False

    def cell(self, row: int, col: int) -> Cell | None:
        return self.cells.get((row, col))

    def column_cells(self, col: int, start_row: int = 1) -> list[Cell]:
        """Non-empty cells of a column, top to bottom."""
        result = []
        for row in range(start_row, self.max_row + 1):
            cell = self.cells.get((row, col))
            if cell is not None and not cell.is_empty:
                result.append(cell)
        return result

    def runs(self, col: int, start_row: int = 1) -> Iterator[list[Cell]]:
        """Contiguous runs of filled cells in a column.

        A run ends at an empty cell: blocks separated by a gap are different
        tables, and comparing formulas across them is meaningless.
        """
        current: list[Cell] = []
        for row in range(start_row, self.max_row + 1):
            cell = self.cells.get((row, col))
            if cell is None or cell.is_empty:
                if current:
                    yield current
                    current = []
            else:
                current.append(cell)
        if current:
            yield current

    def header_row(self) -> int | None:
        """The header row, when one can be recognised.

        The heuristic is deliberately conservative: a header is the first
        non-empty row that is mostly text. When it cannot be recognised, the
        data checks simply run over the whole column.
        """
        for row in range(1, min(self.max_row, 20) + 1):
            row_cells = [self.cells.get((row, c)) for c in range(1, self.max_col + 1)]
            filled = [c for c in row_cells if c is not None and not c.is_empty]
            if len(filled) < 2:
                continue
            texts = sum(1 for c in filled if c.is_text and not c.is_formula)
            if texts / len(filled) < 0.8:
                return None
            below = [self.cells.get((row + 1, c)) for c in range(1, self.max_col + 1)]
            if not [c for c in below if c is not None and not c.is_empty]:
                return None
            return row
        return None

    def header_name(self, col: int) -> str | None:
        row = self.header_row()
        if row is None:
            return None
        cell = self.cells.get((row, col))
        if cell is None or not cell.is_text:
            return None
        return str(cell.value).strip()

    def looks_like_key_column(self, col: int) -> bool:
        name = self.header_name(col)
        return bool(name and _KEY_HEADER_RE.search(name))

    def column_label(self, col: int) -> str:
        name = self.header_name(col)
        letter = get_column_letter(col)
        return f"{letter} '{name}'" if name else letter


@dataclass
class Book:
    path: Path
    sheets: list[Sheet] = field(default_factory=list)
    named_ranges: dict[str, str] = field(default_factory=dict)
    """Workbook-level defined names mapped to their destination. Sheet-local
    names are out of scope for v0.2."""
    has_macros: bool = False
    cache_is_empty: bool = False
    """The workbook was never recalculated: the value cache is empty and
    SF004 cannot see computation errors."""

    def sheet(self, name: str) -> Sheet | None:
        for sheet in self.sheets:
            if sheet.name == name:
                return sheet
        return None


def _read_named_ranges(wb) -> dict[str, str]:
    """Defined names mapped to their destination.

    openpyxl has moved this API around between versions, so the read is
    defensive: a workbook that cannot report its defined names is worth less
    than a crash on every other file.
    """
    result: dict[str, str] = {}
    names = getattr(wb, "defined_names", None)
    if names is None:
        return result
    try:
        items = names.items()
    except AttributeError:  # pragma: no cover - very old openpyxl
        items = ((d.name, d) for d in names.definedName)
    for name, definition in items:
        destination = getattr(definition, "attr_text", None) or getattr(definition, "value", "")
        result[str(name)] = str(destination)
    return result


def load(path: str | Path) -> Book:
    path = Path(path)
    formulas_wb = load_workbook(path, data_only=False)
    values_wb = load_workbook(path, data_only=True)

    book = Book(path=path, has_macros=getattr(formulas_wb, "vba_archive", None) is not None)
    book.named_ranges = _read_named_ranges(formulas_wb)

    formula_cells = 0
    cached_values = 0

    for ws in formulas_wb.worksheets:
        vs = values_wb[ws.title]
        sheet = Sheet(
            name=ws.title,
            max_row=ws.max_row or 0,
            max_col=ws.max_column or 0,
            hidden=ws.sheet_state != "visible",
            merged=[str(r) for r in ws.merged_cells.ranges],
        )
        sheet.hidden_rows = {i for i, d in ws.row_dimensions.items() if d.hidden}
        sheet.hidden_cols = {
            ws[f"{letter}1"].column for letter, d in ws.column_dimensions.items() if d.hidden
        }

        for row in ws.iter_rows():
            for cell in row:
                raw = cell.value
                if raw is None:
                    continue

                formula = None
                if isinstance(raw, str) and raw.startswith("="):
                    formula = raw
                    formula_cells += 1

                cached = vs.cell(row=cell.row, column=cell.column).value
                if formula is not None and cached is not None:
                    cached_values += 1

                value = cached if formula is not None else raw
                if value is None and formula is None:
                    continue

                sheet.cells[(cell.row, cell.column)] = Cell(
                    row=cell.row,
                    col=cell.column,
                    formula=formula,
                    value=value,
                    data_type=cell.data_type,
                    number_format=cell.number_format or "General",
                    is_date=bool(getattr(cell, "is_date", False)),
                )

        book.sheets.append(sheet)

    book.cache_is_empty = formula_cells > 0 and cached_values == 0

    formulas_wb.close()
    values_wb.close()
    return book
