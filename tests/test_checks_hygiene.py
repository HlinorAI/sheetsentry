"""Checks in family SW.

These need workbook features the shared builder does not expose - hidden
rows, merged ranges, defined names - so each test constructs its file with
openpyxl directly.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from sheetsentry import checks, load
from sheetsentry.finding import Severity


@pytest.fixture
def run(tmp_path):
    """Save a workbook built by `prepare` and return findings for one rule."""

    def _run(prepare, rule: str):
        wb = Workbook()
        prepare(wb)
        path = tmp_path / "case.xlsx"
        wb.save(path)
        return [f for f in checks.run(load(path)) if f.rule == rule]

    return _run


def budget(wb, total_row: int = 8) -> None:
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Amount"])
    for index in range(1, total_row - 1):
        ws.append([f"item {index}", index * 100])
    ws[f"B{total_row}"] = f"=SUM(B2:B{total_row - 1})"


class TestSW001HiddenRowsInAggregate:
    def test_hidden_row_inside_a_sum_is_reported(self, run):
        def prepare(wb):
            budget(wb)
            wb.active.row_dimensions[4].hidden = True

        found = run(prepare, "SW001")
        assert len(found) == 1
        assert found[0].severity is Severity.WARNING
        assert found[0].cells == ["B8"]
        assert "hidden rows 4" in found[0].evidence

    def test_no_hidden_rows_is_silent(self, run):
        assert run(budget, "SW001") == []

    def test_hidden_row_outside_the_range_is_ignored(self, run):
        def prepare(wb):
            budget(wb)
            wb.active.row_dimensions[20].hidden = True

        assert run(prepare, "SW001") == []

    def test_hidden_but_empty_row_is_not_reported(self, run):
        # Hiding a spacer row changes nothing about the total.
        def prepare(wb):
            ws = wb.active
            ws.title = "Budget"
            ws.append(["Item", "Amount"])
            for index in range(1, 5):
                ws.append([f"item {index}", index * 100])
            ws["B10"] = "=SUM(B2:B9)"
            ws.row_dimensions[7].hidden = True

        assert run(prepare, "SW001") == []


class TestSW002MergedCells:
    def test_merge_in_the_data_area_is_reported(self, run):
        def prepare(wb):
            ws = wb.active
            ws.title = "Price"
            ws.append(["SKU", "Name", "Price"])
            for index in range(1, 8):
                ws.append([f"A{index}", f"item {index}", index * 10])
            ws.merge_cells("A5:B5")

        found = run(prepare, "SW002")
        assert len(found) == 1
        assert found[0].severity is Severity.WARNING
        assert found[0].evidence == "A5:B5"

    def test_merged_header_is_not_a_finding(self, run):
        def prepare(wb):
            ws = wb.active
            ws.title = "Price"
            ws.append(["SKU", "Name", "Price"])
            for index in range(1, 8):
                ws.append([f"A{index}", f"item {index}", index * 10])
            ws.merge_cells("A1:B1")

        assert run(prepare, "SW002") == []

    def test_clean_sheet_is_silent(self, run):
        def prepare(wb):
            ws = wb.active
            ws.title = "Price"
            ws.append(["SKU", "Price"])
            for index in range(1, 8):
                ws.append([f"A{index}", index])

        assert run(prepare, "SW002") == []


class TestSW003BrokenNamedRange:
    def test_ref_error_in_a_name_is_an_error(self, run):
        def prepare(wb):
            ws = wb.active
            ws.title = "Data"
            ws["A1"] = 1
            wb.defined_names.add(DefinedName("TaxRate", attr_text="Data!#REF!"))

        found = run(prepare, "SW003")
        assert len(found) == 1
        assert found[0].severity is Severity.ERROR
        assert "TaxRate" in found[0].message

    def test_healthy_name_is_silent(self, run):
        def prepare(wb):
            ws = wb.active
            ws.title = "Data"
            ws["A1"] = 1
            wb.defined_names.add(DefinedName("TaxRate", attr_text="Data!$A$1"))

        assert run(prepare, "SW003") == []
